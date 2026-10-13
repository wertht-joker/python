"""简历文件解析：PDF / Word / 纯文本 → 文本"""
import os

SUPPORTED_TYPES = {"pdf", "docx", "txt", "md"}


def file_type(filename: str) -> str:
    return filename.rsplit(".", 1)[-1].lower() if "." in filename else ""


def parse_pdf(path: str) -> str:
    import fitz  # PyMuPDF

    doc = fitz.open(path)
    text = "\n".join(page.get_text() for page in doc)
    doc.close()
    return text


def parse_docx(path: str) -> str:
    from docx import Document

    doc = Document(path)
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    # 表格里的文字（简历常用表格排版）
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text.strip())
    return "\n".join(parts)


def parse_file(path: str, ftype: str | None = None) -> str:
    """按扩展名解析文件为纯文本"""
    ftype = ftype or file_type(path)
    if ftype == "pdf":
        return parse_pdf(path)
    if ftype == "docx":
        return parse_docx(path)
    if ftype in {"txt", "md"}:
        with open(path, "r", encoding="utf-8", errors="ignore") as f:
            return f.read()
    raise ValueError(f"不支持的文件格式: {ftype}")
